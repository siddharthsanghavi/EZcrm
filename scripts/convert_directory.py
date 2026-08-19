"""
Convert the Georgia automation directory spreadsheet into a CSV that EZcrm's
importer understands.

    python scripts/convert_directory.py <path-to-xlsx> [outdir]

Only the "All Companies" sheet is read; the per-type sheets are strict subsets
of it, so importing them too would just create duplicates.
"""

import csv
import sys
from pathlib import Path

import openpyxl

SHEET = "All Companies"

# A club wants two different things from these companies, and which one depends
# on what the company is. Big plants have floors worth walking; the automation
# firms are the ones with a marketing budget. Both are editable after import.
INTEREST_BY_TYPE = {
    "Large Manufacturer": "tour",
    "OEM": "sponsorship;tour",
    "Integrator": "sponsorship;tour",
    "Both": "sponsorship;tour",
    "Machine Shop": "sponsorship",
    "Supplier/Distributor": "sponsorship",
    "Power Automation": "sponsorship",
    "Residential Automation": "sponsorship",
}

COLUMNS = [
    "name", "website", "industry", "status", "interest", "notes",
    "type", "tier", "city", "region", "address", "phone", "employees",
]


def clean(value) -> str:
    if value is None:
        return ""
    text = str(value).strip()
    return "" if text.lower() in ("", "none", "n/a", "-") else text


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__)
        return 1

    src = Path(sys.argv[1])
    outdir = Path(sys.argv[2]) if len(sys.argv) > 2 else Path("data")
    outdir.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if SHEET not in wb.sheetnames:
        print(f"No '{SHEET}' sheet found. Sheets: {wb.sheetnames}")
        return 1

    rows = list(wb[SHEET].iter_rows(values_only=True))
    header = [clean(h) for h in rows[0]]

    seen: set[str] = set()
    out: list[dict[str, str]] = []
    skipped_dupes = 0

    for raw in rows[1:]:
        row = dict(zip(header, raw))
        name = clean(row.get("Company"))
        if not name:
            continue

        key = name.lower()
        if key in seen:
            skipped_dupes += 1
            continue
        seen.add(key)

        ctype = clean(row.get("Type"))

        # Everything that doesn't map to a column but is still worth reading
        # lands in notes, one fact per line so the record stays scannable.
        note_parts = []
        if clean(row.get("Notes")):
            note_parts.append(clean(row["Notes"]))
        if clean(row.get("PLC_Oriented")):
            note_parts.append(f"PLC-oriented: {clean(row['PLC_Oriented'])}")
        if clean(row.get("Captive_Builder")):
            note_parts.append(f"Captive builder: {clean(row['Captive_Builder'])}")

        out.append({
            "name": name,
            "website": clean(row.get("Website")),
            "industry": clean(row.get("Specialty")),
            "status": "prospect",
            "interest": INTEREST_BY_TYPE.get(ctype, "sponsorship"),
            "notes": "\n".join(note_parts),
            "type": ctype,
            "tier": clean(row.get("Confidence")),
            "city": clean(row.get("City")),
            "region": clean(row.get("Region")),
            "address": clean(row.get("Address")),
            "phone": clean(row.get("Phone")),
            "employees": clean(row.get("Employees")),
        })

    def write(path: Path, records: list[dict[str, str]]) -> None:
        with path.open("w", newline="", encoding="utf-8") as fh:
            writer = csv.DictWriter(fh, fieldnames=COLUMNS)
            writer.writeheader()
            writer.writerows(records)
        print(f"  {len(records):5d}  {path}")

    # The full list, plus a starter file. 1295 companies is more than a club can
    # work; the tiers are the author's own confidence ranking, so Tier 1 and 2
    # are the ones actually worth calling first.
    priority = [r for r in out if r["tier"] in ("Tier 1", "Tier 2")]

    print(f"\nParsed {len(out)} companies ({skipped_dupes} duplicate name(s) dropped)\n")
    write(outdir / "companies-all.csv", out)
    write(outdir / "companies-priority.csv", priority)

    print("\nBy tier:")
    for tier in ("Tier 1", "Tier 2", "Tier 3", "Reference"):
        print(f"  {sum(1 for r in out if r['tier'] == tier):5d}  {tier}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
